from __future__ import annotations

import hashlib
import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from ..monitor import diff_against_baseline
from ..state.database import Database

REPORT_VERSION = "1.0"


def report_output_dir(task_root: str | Path, task_id: str) -> Path:
    """C1: one directory per task under reports/<task-id>/ - consecutive
    `report` runs of different tasks no longer overwrite each other.
    Legacy flat reports/ files are left in place, never migrated."""
    return Path(task_root) / "reports" / task_id


def compute_integrity_status(
    task_status: str,
    *,
    gaps: list[dict],
    unknown_actions: int,
    failed_actions: int,
    pending_actions: int,
    total_actions: int = 0,
) -> str:
    if task_status in {"failed"}:
        return "failed"
    if task_status in {"paused", "blocked", "cancelled", "needs_review"}:
        return "partial"
    if total_actions == 0:
        # 真实 recon 不写 action 行（plan/run 才写）：action 维度不得把
        # 无 action 行的任务拖成 unknown。缺口仍诚实降级 partial。
        return "partial" if gaps else "complete"
    if failed_actions and not gaps:
        return "partial"
    if unknown_actions or pending_actions:
        return "unknown" if not gaps else "partial"
    if gaps:
        return "partial"
    return "complete"


def build_report(db: Database, task_id: str, output_dir: str | Path,
                 *, reveal_emails: bool = False,
                 include_xlsx: bool = False) -> dict:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    task = next(db.rows("SELECT * FROM task WHERE task_id=?", (task_id,)), None)
    if task is None:
        raise ValueError("task_not_found")
    assets = [dict(row) for row in db.rows("SELECT * FROM asset WHERE task_id=?", (task_id,))]
    asset_relations = [dict(row) for row in db.rows(
        "SELECT * FROM asset_relation WHERE task_id=? ORDER BY rowid", (task_id,),
    )]
    asset_conflicts = [dict(row) for row in db.rows(
        "SELECT * FROM asset_conflict WHERE task_id=? ORDER BY rowid", (task_id,),
    )]
    observations = [dict(row) for row in db.rows("SELECT * FROM observation WHERE task_id=?", (task_id,))]
    for row in observations:
        # P04: expose the payload as a real object next to the raw string the
        # runners wrote - the string field stays verbatim for existing
        # consumers, the object spares every downstream reader a json.loads
        raw = row.get("payload_json")
        if raw:
            try:
                row["payload"] = json.loads(raw)
            except (TypeError, ValueError):
                row["payload"] = None
    gaps = [dict(row) for row in db.rows("SELECT * FROM gap WHERE task_id=?", (task_id,))]
    evidence = [dict(row) for row in db.rows("SELECT * FROM evidence WHERE task_id=?", (task_id,))]
    events = [dict(row) for row in db.rows("SELECT * FROM event WHERE task_id=? ORDER BY created_at", (task_id,))]
    checkpoints = [dict(row) for row in db.rows("SELECT * FROM checkpoint WHERE task_id=? ORDER BY created_at", (task_id,))]
    fofa_batches = [dict(row) for row in db.rows(
        "SELECT * FROM fofa_query_batch WHERE task_id=? ORDER BY batch_index, batch_id", (task_id,)
    )]
    fofa_plans = [dict(row) for row in db.rows(
        "SELECT * FROM fofa_query_plan WHERE task_id=? ORDER BY created_at", (task_id,)
    )]
    budgets = [dict(row) for row in db.rows("SELECT * FROM budget WHERE task_id=?", (task_id,))]
    adapter_rows = [dict(row) for row in db.rows(
        "SELECT * FROM adapter_registry ORDER BY adapter_name")]
    scope_rows = [dict(row) for row in db.rows(
        "SELECT scope_id,scope_version,payload_json,confirmed,created_at FROM scope_snapshot "
        "WHERE task_id=? ORDER BY created_at", (task_id,))]
    for row in scope_rows:
        try:
            row["payload"] = json.loads(row.pop("payload_json"))
        except json.JSONDecodeError:
            row["payload"] = None
    unknown_actions = sum(1 for row in db.rows(
        "SELECT status FROM action WHERE task_id=? AND status='unknown'", (task_id,)
    ))
    failed_actions = sum(1 for row in db.rows(
        "SELECT status FROM action WHERE task_id=? AND status='failed'", (task_id,)
    ))
    pending_actions = sum(1 for row in db.rows(
        "SELECT status FROM action WHERE task_id=? AND status IN ('planned','started')", (task_id,)
    ))
    total_actions = next(db.rows(
        "SELECT COUNT(*) AS n FROM action WHERE task_id=?", (task_id,)
    ), {"n": 0})["n"]
    integrity_status = compute_integrity_status(
        task["status"],
        gaps=gaps,
        unknown_actions=unknown_actions,
        failed_actions=failed_actions,
        pending_actions=pending_actions,
        total_actions=total_actions,
    )
    stop_events = [
        {"event_type": row["event_type"], "payload": json.loads(row["payload_json"]),
         "created_at": row["created_at"]}
        for row in events if row["event_type"] in {"task_paused", "task_cancelled"}
    ]
    review_queue = [
        {**row, "kind": "architecture_block" if row["reason"].startswith("architecture_blocked") else "generic"}
        for row in gaps
    ]
    observations_by_source: dict[str, int] = {}
    for row in observations:
        observations_by_source[row["source"]] = observations_by_source.get(row["source"], 0) + 1
    assets_by_type: dict[str, int] = {}
    for row in assets:
        assets_by_type[row["asset_type"]] = assets_by_type.get(row["asset_type"], 0) + 1
    # retrospective P0-2: deduplicate certificate observations per
    # (common_name, issuer, not_after) - re-observations are legitimate but
    # must not read as extra certificates; raw rows stay in observations
    certificates: list[dict] = []
    seen_certs: dict[tuple, dict] = {}
    for row in observations:
        if row["source"] != "ct_log":
            continue
        try:
            payload = json.loads(row["payload_json"])
        except json.JSONDecodeError:
            continue
        for cert in payload.get("certificates", []):
            key = (cert.get("common_name"), cert.get("issuer"), cert.get("not_after"))
            if key not in seen_certs:
                entry = {**cert, "observation_count": 0}
                seen_certs[key] = entry
                certificates.append(entry)
            seen_certs[key]["observation_count"] += 1
    browser_pages = [dict(row) for row in observations if row["source"] == "playwright_mcp"]
    try:
        browser_network = [
            entry
            for page in browser_pages
            for entry in json.loads(page["payload_json"]).get("network_summary", [])
        ]
    except json.JSONDecodeError:
        browser_network = []
    diff = diff_against_baseline(db, task_id, task["scope_version"])
    snapshot_diff = {
        "baseline_scope_version": diff.baseline_scope_version,
        "scope_version_match": diff.scope_version_match,
        "added": diff.added,
        "disappeared": diff.disappeared,
        "changed": diff.changed,
        "reappeared": diff.reappeared,
        "unchanged": diff.unchanged,
    }
    report = {
        "report_version": REPORT_VERSION,
        "task": dict(task),
        "integrity_status": integrity_status,
        "summary": {
            "task_id": task_id,
            "status": task["status"],
            "scope_version": task["scope_version"],
            "integrity_status": integrity_status,
            "asset_count": len(assets),
            "assets_by_type": assets_by_type,
            "observation_count": len(observations),
            "observations_by_source": observations_by_source,
            "evidence_count": len(evidence),
            "unresolved_conflicts": len(asset_conflicts),
            "review_queue_count": len(review_queue),
            "gap_count": len(gaps),
            "budget_lines": len(budgets),
            # C4: distinguish "no budget configured" (opt-in mechanism, no
            # rows) from a real zero-consumption budget
            "budget_configured": bool(budgets),
            "budget_consumed_units": sum(row["consumed_units"] for row in budgets),
            "planned_actions_pending": pending_actions,
            "integrity_unknown_reason": (
                f"{pending_actions} planned action(s) were never executed - direct "
                "runner commands bypass the plan loop" if integrity_status == "unknown"
                and pending_actions else None),
            "stop_events": stop_events,
            "generated_at": datetime.now(timezone.utc).isoformat(),
        },
        "scope_authorization": scope_rows,
        "adapter_status": adapter_rows,
        "budgets": budgets,
        "coverage": {
            "observations_by_source": observations_by_source,
            "assets_by_type": assets_by_type,
            "fofa_batches_completed": sum(row["status"] == "completed" for row in fofa_batches),
            "fofa_batches_failed": sum(row["status"] == "failed" for row in fofa_batches),
        },
        "quality": {
            "integrity_status": integrity_status,
            "unresolved_conflicts": len(asset_conflicts),
            "review_queue": review_queue,
            "gaps": gaps,
        },
        "merge_summary": {
            "asset_count": len(assets),
            "confirmed_assets": sum(row["scope_status"] == "confirmed" for row in assets),
            "candidate_assets": sum(row["scope_status"] == "candidate" for row in assets),
            "third_party_assets": sum(row["architecture_status"] == "third_party" for row in assets),
            "relation_count": len(asset_relations),
            "confirmed_relations": sum(row["status"] == "confirmed" for row in asset_relations),
            "unresolved_conflicts": len(asset_conflicts),
            "sources": sorted({item["source"] for row in assets
                               for item in json.loads(row.get("sources_json") or "[]")}),
        },
        "assets": assets,
        "services": [row for row in assets if row["asset_type"] == "service"],
        "domains_certificates": [row for row in assets if row["asset_type"] == "domain"],
        "certificates": certificates,
        "web_stack": [
            {"asset_id": row["asset_id"], "normalized_value": row["normalized_value"],
             "architecture_status": row["architecture_status"],
             "evidence": json.loads(row.get("architecture_json") or "[]")}
            for row in assets if row.get("architecture_json")
        ],
        "asset_relations": asset_relations,
        "asset_conflicts": asset_conflicts,
        "browser": {
            "pages": browser_pages,
            "network": browser_network,
        },
        "observations": observations,
        "evidence": evidence,
        "events": events,
        "checkpoints": checkpoints,
        "fofa": {
            "plans": fofa_plans,
            "batches": fofa_batches,
            "coverage": {
                "planned": len(fofa_plans),
                "batches": len(fofa_batches),
                "completed_batches": sum(row["status"] == "completed" for row in fofa_batches),
                "failed_batches": sum(row["status"] == "failed" for row in fofa_batches),
                "observations": sum(row["result_count"] for row in fofa_batches),
                "estimated_fpoints": sum(row["estimated_fpoints"] for row in fofa_batches),
                "consumed_fpoints": sum(row["consumed_fpoints"] for row in fofa_batches),
            },
        },
        "snapshot_diff": snapshot_diff,
        "recovery_summary": {
            "last_checkpoint": checkpoints[-1]["checkpoint_id"] if checkpoints else None,
            "recoverable": bool(checkpoints),
            "unknown_actions": unknown_actions,
            "failed_actions": failed_actions,
            "pending_actions": pending_actions,
        },
        "export": {},
    }
    if reveal_emails:
        # C3: explicit opt-in - substitute the redaction_map plaintexts back
        # into every deliverable; the default keeps placeholders only
        _reveal_emails(db, task_id, report)
    json_path = output / "result.json"
    json_path.write_text(json.dumps(report, ensure_ascii=True, indent=2), encoding="utf-8")
    # 2026-10-05 (operator decision): result.json is THE machine-readable
    # intermediate - the xlsx renderings (result.xlsx / analyst_report.xlsx)
    # are opt-in via include_xlsx for humans inspecting a draft. The final
    # delivered workbook is built from this JSON by
    # scripts/build_deliverable_xlsx.py, so the default report leaves no xlsx
    # intermediates behind.
    if include_xlsx:
        report["export"] = _export_xlsx(output, report, task_id, assets, scope_rows,
                                       asset_relations, asset_conflicts, browser_pages,
                                       browser_network, review_queue, checkpoints)
        # analyst-facing delivery in the company template; failure isolated from
        # the machine report exactly like the XLSX path (case-08 semantics)
        from .analyst_report import build_analyst_report

        analyst_meta = build_analyst_report(db, task_id, report, output)
        report["export"]["analyst_report"] = analyst_meta
        # export metadata is appended to the JSON after the XLSX attempt; case-08:
        # a failing XLSX never invalidates the already-written JSON main result
        json_path.write_text(json.dumps(report, ensure_ascii=True, indent=2), encoding="utf-8")
    return report


def _reveal_emails(db: Database, task_id: str, report: dict) -> None:
    """C3: deep-replace local redaction_map placeholders with the plaintext
    emails. No map entries -> nothing changes."""
    pairs = [(row["placeholder"], row["plaintext"]) for row in db.rows(
        "SELECT placeholder,plaintext FROM redaction_map WHERE task_id=?", (task_id,))]
    if not pairs:
        return

    def walk(node):
        if isinstance(node, dict):
            return {key: walk(value) for key, value in node.items()}
        if isinstance(node, list):
            return [walk(item) for item in node]
        if isinstance(node, tuple):
            return tuple(walk(item) for item in node)
        if isinstance(node, str):
            for placeholder, plaintext in pairs:
                if placeholder in node:
                    node = node.replace(placeholder, plaintext)
        return node

    for key in list(report):
        report[key] = walk(report[key])


def _export_xlsx(output: Path, report: dict, task_id: str, assets, scope_rows,
                 asset_relations, asset_conflicts, browser_pages, browser_network,
                 review_queue, checkpoints) -> dict:
    """Atomic XLSX export; failure keeps the JSON main result authoritative."""
    started_at = datetime.now(timezone.utc).isoformat()
    xlsx_path = output / "result.xlsx"
    try:
        _build_workbook(xlsx_path, report, task_id, assets, scope_rows, asset_relations,
                        asset_conflicts, browser_pages, browser_network, review_queue,
                        checkpoints)
        data = xlsx_path.read_bytes()
        return {
            "json": str(xlsx_path.parent / "result.json"),
            "xlsx": str(xlsx_path),
            "xlsx_sha256": hashlib.sha256(data).hexdigest(),
            "xlsx_size": len(data),
            "xlsx_status": "exported",
            "report_version": report["report_version"],
            "asset_count_consistent": len(assets) == report["summary"]["asset_count"],
            "integrity_status_consistent": report["summary"]["integrity_status"],
            "exported_at": started_at,
        }
    except Exception as exc:
        # keep JSON authoritative; XLSX stays absent or partial-unpublished
        return {
            "json": str(xlsx_path.parent / "result.json"),
            "xlsx": str(xlsx_path),
            "xlsx_status": "failed",
            "xlsx_error": f"{type(exc).__name__}: {exc}",
            "report_version": report["report_version"],
            "exported_at": started_at,
        }


def _build_workbook(xlsx_path: Path, report: dict, task_id: str, assets, scope_rows,
                    asset_relations, asset_conflicts, browser_pages, browser_network,
                    review_queue, checkpoints) -> None:
    workbook = Workbook()
    workbook.active.title = "报告摘要"
    workbook["报告摘要"].append(["task_id", "report_version", "integrity_status",
                                 "asset_count", "observation_count", "unresolved_conflicts"])
    summary = report["summary"]
    workbook["报告摘要"].append([
        task_id, report["report_version"], summary["integrity_status"],
        summary["asset_count"], summary["observation_count"],
        summary["unresolved_conflicts"],
    ])
    sheet = workbook.create_sheet("资产总表")
    if assets:
        sheet.append(list(assets[0].keys()))
        for row in assets:
            sheet.append([_excel_safe(row[key]) for key in assets[0]])
    scope_sheet = workbook.create_sheet("范围授权")
    for row in scope_rows:
        scope_sheet.append([_excel_safe(row[key]) for key in
                            ("scope_id", "scope_version", "confirmed", "created_at")])
        scope_sheet.append([_excel_safe(json.dumps(row.get("payload") or {}, sort_keys=True))])
    services = report["services"]
    service_sheet = workbook.create_sheet("服务")
    if services:
        service_sheet.append(list(services[0].keys()))
        for row in services:
            service_sheet.append([_excel_safe(row[key]) for key in services[0]])
    domain_sheet = workbook.create_sheet("域名证书")
    for row in report["domains_certificates"]:
        domain_sheet.append([_excel_safe(row["normalized_value"]), _excel_safe(row["scope_status"]),
                             _excel_safe(row["confidence"]), _excel_safe(row["current_status"])])
    web_sheet = workbook.create_sheet("Web技术")
    for row in report["web_stack"]:
        web_sheet.append([_excel_safe(row["normalized_value"]),
                          _excel_safe(row["architecture_status"]),
                          _excel_safe(json.dumps(row["evidence"], sort_keys=True))])
    relations_sheet = workbook.create_sheet("架构关系")
    if asset_relations:
        relations_sheet.append(list(asset_relations[0].keys()))
        for row in asset_relations:
            relations_sheet.append([_excel_safe(row[key]) for key in asset_relations[0]])
    page_sheet = workbook.create_sheet("浏览器页面")
    for row in browser_pages:
        page_sheet.append([_excel_safe(row["target"]), _excel_safe(row["payload_json"][:1000])])
    network_sheet = workbook.create_sheet("浏览器网络")
    for entry in browser_network:
        network_sheet.append([_excel_safe(json.dumps(entry, sort_keys=True))])
    conflicts_sheet = workbook.create_sheet("归并冲突")
    conflicts_sheet.append(["conflict_id", "asset_id", "attribute", "values_json", "status", "created_at"])
    for row in asset_conflicts:
        conflicts_sheet.append([
            _excel_safe(row["conflict_id"]), _excel_safe(row["asset_id"]),
            _excel_safe(row["attribute"]), _excel_safe(row["values_json"]),
            _excel_safe(row["status"]), _excel_safe(row["created_at"]),
        ])
    workbook.create_sheet("信息缺口")
    for row in report["quality"]["gaps"]:
        workbook["信息缺口"].append([_excel_safe(value) for value in row.values()])
    review_sheet = workbook.create_sheet("review_queue")
    review_sheet.append(["gap_id", "kind", "priority", "reason", "status"])
    for row in review_queue:
        review_sheet.append([_excel_safe(row["gap_id"]), _excel_safe(row["kind"]),
                             _excel_safe(row["priority"]), _excel_safe(row["reason"]),
                             _excel_safe(row["status"])])
    evidence_sheet = workbook.create_sheet("证据文件")
    for row in report["evidence"]:
        evidence_sheet.append([_excel_safe(value) for value in row.values()])
    log_sheet = workbook.create_sheet("执行日志")
    for row in report["events"]:
        log_sheet.append([_excel_safe(value) for value in row.values()])
    recovery_sheet = workbook.create_sheet("checkpoint_recovery")
    recovery_sheet.append(["checkpoint_id", "phase", "status", "created_at"])
    for row in checkpoints:
        recovery_sheet.append([_excel_safe(row["checkpoint_id"]), _excel_safe(row["phase"]),
                               _excel_safe(row["status"]), _excel_safe(row["created_at"])])
    recovery_sheet.append([])
    recovery_sheet.append(["last_checkpoint", "recoverable", "unknown_actions",
                           "failed_actions", "pending_actions"])
    recovery = report["recovery_summary"]
    recovery_sheet.append([_excel_safe(recovery["last_checkpoint"]), _excel_safe(recovery["recoverable"]),
                           _excel_safe(recovery["unknown_actions"]), _excel_safe(recovery["failed_actions"]),
                           _excel_safe(recovery["pending_actions"])])
    diff_sheet = workbook.create_sheet("快照差异")
    diff = report["snapshot_diff"]
    diff_sheet.append(["baseline_scope_version", "scope_version_match"])
    diff_sheet.append([_excel_safe(diff["baseline_scope_version"]),
                       _excel_safe(diff["scope_version_match"])])
    diff_sheet.append([])
    for section in ("added", "disappeared", "changed", "reappeared"):
        diff_sheet.append([section])
        for row in diff[section]:
            diff_sheet.append([_excel_safe(json.dumps(row, sort_keys=True))])
    fofa_sheet = workbook.create_sheet("FOFA查询")
    fofa_sheet.append(["query_plan_count", "batch_count", "completed_batches", "failed_batches",
                       "observations", "estimated_fpoints", "consumed_fpoints"])
    coverage = report["fofa"]["coverage"]
    fofa_sheet.append([coverage["planned"], coverage["batches"], coverage["completed_batches"],
                       coverage["failed_batches"], coverage["observations"],
                       coverage["estimated_fpoints"], coverage["consumed_fpoints"]])
    # atomic publish: the workbook is built in a temp file then moved into place
    handle = tempfile.NamedTemporaryFile(prefix="result-", suffix=".xlsx",
                                         dir=str(xlsx_path.parent), delete=False)
    handle.close()
    temp_path = Path(handle.name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, xlsx_path)
    except Exception:
        try:
            temp_path.unlink(missing_ok=True)
        finally:
            raise


def _excel_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
