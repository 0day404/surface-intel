import asyncio
import json

from openpyxl import load_workbook

from surface_intel.cli import main
from surface_intel.coordinator import run_offline_task
from surface_intel.reporting.export import build_report
from surface_intel.state.database import Database


def test_offline_flow_persists_observations_and_rebuilds_reports(tmp_path, capsys):
    main(["init", str(tmp_path)])
    capsys.readouterr()
    scope = tmp_path / "scope.json"
    scope.write_text(json.dumps({"authorized_domains": ["example.com"], "confirmed": True}), encoding="utf-8")
    main(["create", "--task-root", str(tmp_path), "--scope", str(scope)])
    task_id = json.loads(capsys.readouterr().out)["task_id"]
    main(["plan", "--task-root", str(tmp_path), task_id])
    capsys.readouterr()
    result = asyncio.run(run_offline_task(Database(tmp_path / "task.sqlite"), task_id, [
        {"observations": [{"host": "edge.example.com"}], "has_more": False}
    ]))
    assert result["assets"] == 1
    report = build_report(Database(tmp_path / "task.sqlite"), task_id, tmp_path / "reports")
    assert report["integrity_status"] == "partial"
    assert len(report["assets"]) == 1
    assert load_workbook(tmp_path / "reports" / "result.xlsx").sheetnames[:2] == ["报告摘要", "资产总表"]
