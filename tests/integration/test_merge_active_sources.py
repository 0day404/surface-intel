"""Active-observation sources (safe_http / safe_port) in the merge layer (B1).

The extractor whitelist used to cover the passive sources only, so active
verification never reached the asset model and the analyst 站点/服务 sheets
stayed empty for pure active tasks. These tests pin the new behavior:
attribute refresh on observed responses, service assets for IP hosts,
scope-anchor guard against forged observations, idempotent re-merge.
"""
import json

from openpyxl import load_workbook

from surface_intel.merger import merge_task_observations
from surface_intel.policy.scope import Scope
from surface_intel.reporting.export import build_report
from surface_intel.state.database import Database


def _setup(tmp_path):
    scope = Scope(authorized_domains=("example.com", "*.example.com"),
                  authorized_ips=("203.0.113.5",), confirmed=True,
                  scope_version="scope-v1")
    db = Database(tmp_path / "task.sqlite")
    db.initialize()
    with db.connect() as connection:
        connection.execute(
            "INSERT INTO task(task_id,status,schema_version,scope_version) VALUES('t','planned','1.0',?)",
            (scope.scope_version,),
        )
    return db, scope


def _observation(db, task_id, source, payload, observed_at="2026-10-04T10:00:00+00:00"):
    payload = dict(payload)
    payload.setdefault("observed_at", observed_at)
    with db.connect() as connection:
        connection.execute(
            """INSERT INTO observation(observation_id,task_id,source,target,payload_json)
               VALUES(?,?,?,?,?)""",
            (f"obs-{source}-{abs(hash(json.dumps(payload, sort_keys=True))) % 10**12}",
             task_id, source, str(payload.get("target", "")),
             json.dumps(payload, ensure_ascii=True, sort_keys=True)),
        )


def _rows(db, query, params=()):
    return [dict(row) for row in db.rows(query, params)]


def _assets(db):
    return {(row["asset_type"], row["normalized_value"]): row
            for row in _rows(db, "SELECT * FROM asset")}


def test_safe_http_observation_refreshes_host_asset_attributes(tmp_path):
    db, scope = _setup(tmp_path)
    _observation(db, "t", "dns_resolver", {
        "target": "www.example.com", "record_type": "A", "records": ["203.0.113.5"],
        "status": "observed",
    })
    _observation(db, "t", "safe_http", {
        "target": "https://www.example.com/", "transport_status": "observed",
        "status_code": 200, "final_url": "https://www.example.com/",
        "headers": {"server": "nginx", "content-type": "text/html"},
        "content_type": "text/html", "redirect_chain": [],
    })
    summary = merge_task_observations(db, "t", scope)
    assert summary.observations_processed == 2
    assets = _assets(db)
    domain = assets[("domain", "www.example.com")]
    assert domain["scope_status"] == "confirmed"
    attributes = json.loads(domain["attributes_json"])
    assert attributes["http_status"]["value"] == "200"
    assert attributes["http_server"]["value"] == "nginx"
    assert attributes["http_final_url"]["value"] == "https://www.example.com/"
    assert attributes["http_content_type"]["value"] == "text/html"
    # the observation refreshes facts, it does not mint new assets
    assert ("domain", "https://www.example.com") not in assets


def test_failed_safe_http_transport_creates_no_assets(tmp_path):
    db, scope = _setup(tmp_path)
    _observation(db, "t", "safe_http", {
        "target": "https://www.example.com/", "transport_status": "tls_error",
        "redirect_chain": [], "error_type": "SSL Cert Verification Error",
    })
    summary = merge_task_observations(db, "t", scope)
    assert summary.observations_processed == 1
    assert _assets(db) == {}


def test_safe_port_open_ip_creates_service_asset_and_relation(tmp_path):
    db, scope = _setup(tmp_path)
    _observation(db, "t", "safe_port", {
        "target": "203.0.113.5", "port": 443, "transport_status": "open",
        "tls_sha256": "aa" * 32,
    })
    summary = merge_task_observations(db, "t", scope)
    assets = _assets(db)
    service = assets[("service", "203.0.113.5:443/https")]
    assert service["scope_status"] == "confirmed"
    assert service["port"] == 443 and service["protocol"] == "https"
    assert assets[("ip", "203.0.113.5")]["scope_status"] == "confirmed"
    relations = _rows(db, "SELECT * FROM asset_relation")
    assert len(relations) == 1
    assert relations[0]["relation_type"] == "service_ip"
    assert relations[0]["relation_strength"] == "strong"
    assert relations[0]["status"] == "confirmed"


def test_safe_port_without_tls_evidence_stays_unknown_protocol(tmp_path):
    db, scope = _setup(tmp_path)
    _observation(db, "t", "safe_port", {
        "target": "203.0.113.5", "port": 8080, "transport_status": "open",
    })
    merge_task_observations(db, "t", scope)
    assert ("service", "203.0.113.5:8080/unknown") in _assets(db)


def test_safe_port_closed_creates_nothing(tmp_path):
    db, scope = _setup(tmp_path)
    _observation(db, "t", "safe_port", {
        "target": "203.0.113.5", "port": 443, "transport_status": "closed",
    })
    summary = merge_task_observations(db, "t", scope)
    assert summary.assets_created == 0
    assert _assets(db) == {}


def test_safe_port_domain_target_records_attribute_not_service(tmp_path):
    """Service assets are IP-keyed by normalization; a domain target records
    the outcome as a host attribute plus an auditable skip entry."""
    db, scope = _setup(tmp_path)
    _observation(db, "t", "safe_port", {
        "target": "www.example.com", "port": 8443, "transport_status": "open",
    })
    summary = merge_task_observations(db, "t", scope)
    assets = _assets(db)
    assert not any(key[0] == "service" for key in assets)
    attributes = json.loads(assets[("domain", "www.example.com")]["attributes_json"])
    assert attributes["open_port_8443"]["value"] == "unknown"
    assert any(item["reason"] == "service_requires_ip_host" for item in summary.skipped)


def test_forged_active_observation_cannot_self_confirm(tmp_path):
    """The anchor guard holds for active sources: a host outside the scope
    document stays a candidate no matter what the observation claims."""
    db, scope = _setup(tmp_path)
    _observation(db, "t", "safe_http", {
        "target": "https://other.example.net/", "transport_status": "observed",
        "status_code": 200, "redirect_chain": [],
    })
    merge_task_observations(db, "t", scope)
    assets = _assets(db)
    assert assets[("domain", "other.example.net")]["scope_status"] == "candidate"
    # the observed facts themselves are still recorded on the candidate
    attributes = json.loads(assets[("domain", "other.example.net")]["attributes_json"])
    assert attributes["http_status"]["value"] == "200"


def test_remerge_keeps_attributes_stable_and_idempotent(tmp_path):
    db, scope = _setup(tmp_path)
    _observation(db, "t", "safe_http", {
        "target": "https://www.example.com/", "transport_status": "observed",
        "status_code": 200, "redirect_chain": [],
    })
    first = merge_task_observations(db, "t", scope)
    second = merge_task_observations(db, "t", scope)
    assert second.assets_created == 0
    assert second.assets_updated == first.assets_created
    before = _assets(db)[("domain", "www.example.com")]["attributes_json"]
    merge_task_observations(db, "t", scope)
    after = _assets(db)[("domain", "www.example.com")]["attributes_json"]
    assert before == after


def test_conflict_attributes_never_enter_attributes_json(tmp_path):
    """Conflict-bearing values live in asset_conflict rows exclusively - no
    double representation in the attribute column."""
    db, scope = _setup(tmp_path)
    _observation(db, "t", "rdap_lookup", {
        "target": "203.0.113.5", "target_kind": "ip", "status": "observed",
        "record": {"organizations": ["Org A Ltd"]}, "source_protocol": "rdap",
    }, observed_at="2026-10-04T10:00:00+00:00")
    _observation(db, "t", "rdap_lookup", {
        "target": "203.0.113.5", "target_kind": "ip", "status": "observed",
        "record": {"organizations": ["Org B Inc"]}, "source_protocol": "rdap",
    }, observed_at="2026-10-04T11:00:00+00:00")
    summary = merge_task_observations(db, "t", scope)
    assert summary.conflicts_created == 1
    ip_asset = _assets(db)[("ip", "203.0.113.5")]
    assert json.loads(ip_asset["attributes_json"]) == {}


def test_analyst_site_and_service_sheets_reflect_active_sources(tmp_path):
    db, scope = _setup(tmp_path)
    with db.connect() as connection:
        connection.execute(
            "UPDATE task SET status='completed' WHERE task_id='t'")
    _observation(db, "t", "safe_http", {
        "target": "https://www.example.com/", "transport_status": "observed",
        "status_code": 200, "final_url": "https://www.example.com/",
        "redirect_chain": [],
    })
    _observation(db, "t", "safe_port", {
        "target": "203.0.113.5", "port": 443, "transport_status": "open",
        "tls_sha256": "aa" * 32,
    })
    merge_task_observations(db, "t", scope)
    build_report(db, "t", tmp_path / "reports")
    workbook = load_workbook(tmp_path / "reports" / "analyst_report.xlsx")
    sites = [row for row in workbook["站点"].iter_rows(values_only=True)][1:]
    assert ["https://www.example.com/", "HTTP 200", "safe_http 主动观察"] in [list(row) for row in sites]
    services = [row for row in workbook["服务"].iter_rows(values_only=True)][1:]
    assert [list(row)[:3] for row in services] == [["203.0.113.5", "443", "HTTPS"]]
