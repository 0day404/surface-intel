"""Analyst-report generation in the company delivery template.

The analyst workbook is a second deliverable shaped for human review:
Chinese analytic sheets (归属档/用途/依据), placeholder sections for data the
pipeline does not collect, and the machine result.xlsx stays untouched.
"""
import json

from openpyxl import load_workbook

from surface_intel.merger import merge_task_observations
from surface_intel.policy.scope import Scope
from surface_intel.reporting.export import build_report
from surface_intel.state.database import Database


def _seed(tmp_path):
    scope = Scope(authorized_domains=("example.com", "*.example.com"),
                  authorized_ips=("203.0.113.5",), confirmed=True,
                  scope_version="scope-v1")
    db = Database(tmp_path / "task.sqlite")
    db.initialize()
    with db.connect() as connection:
        connection.execute(
            "INSERT INTO task(task_id,status,schema_version,scope_version) VALUES('t','completed','1.0',?)",
            (scope.scope_version,))
        records = [
            {"source": "dns_resolver", "target": "www.example.com", "record_type": "A",
             "records": ["203.0.113.5"], "status": "observed"},
            {"source": "fofa_mcp", "host": "www.example.com", "ip": "203.0.113.5",
             "port": 443, "protocol": "https", "title": "Example 登录",
             "product": "nginx", "status": "observed"},
            {"source": "ct_log", "target": "example.com", "status": "observed",
             "certificates": [{"common_name": "www.example.com",
                               "sans": ["www.example.com", "app.example.com"],
                               "issuer": "Sectigo", "serial": "aa",
                               "not_before": "n", "not_after": "x",
                               "first_seen": "f", "last_seen": "l", "log_id": 1}]},
            {"source": "rdap_lookup", "target": "203.0.113.5", "target_kind": "ip",
             "status": "observed", "record": {"organizations": ["Example Org Ltd"],
                                              "country": "CN"}},
        ]
        for index, payload in enumerate(records):
            payload = dict(payload)
            payload.setdefault("observed_at", f"2026-10-04T10:0{index}:00+00:00")
            connection.execute(
                """INSERT INTO observation(observation_id,task_id,source,target,payload_json)
                   VALUES(?,?,?,?,?)""",
                (f"obs-{index}", "t", payload.pop("source"), str(payload.get("target", "")),
                 json.dumps(payload, sort_keys=True)),
            )
    merge_task_observations(db, "t", scope)
    return db


def test_analyst_report_follows_company_template(tmp_path):
    db = _seed(tmp_path)
    report = build_report(db, "t", tmp_path / "reports", include_xlsx=True)
    meta = report["export"]["analyst_report"]
    assert meta["status"] == "exported"

    workbook = load_workbook(tmp_path / "reports" / "analyst_report.xlsx")
    # company template sheets, in delivery order
    assert workbook.sheetnames == [
        "报告摘要", "域名资产", "IP资产", "站点", "服务", "组件", "登录入口",
        "APP资产", "微信公众号资产", "小程序", "暗网泄露", "信息泄露", "漏洞汇总",
    ]

    domains = list(workbook["域名资产"].iter_rows(values_only=True))
    assert domains[0] == ("域名资产", "归属档", "用途", "已发现子域", "归属依据", "数据来源")
    values = {row[0]: row for row in domains[1:]}
    assert "example.com" in values
    assert values["example.com"][1].startswith("A-")  # confirmed scope tag
    # CT SAN became a discovered-subdomain entry under the apex
    assert "app.example.com" in values["example.com"][3]

    ip_rows = {row[0]: row for row in workbook["IP资产"].iter_rows(values_only=True)
               if row[0] not in (None, "IP")}
    assert ip_rows["203.0.113.5"][2] == "CN"  # RDAP country surfaced
    assert "www.example.com" in ip_rows["203.0.113.5"][3]  # carried assets

    sites = [row[0] for row in workbook["站点"].iter_rows(values_only=True)][1:]
    assert any("www.example.com" in site for site in sites)

    logins = [row for row in workbook["登录入口"].iter_rows(values_only=True)][1:]
    assert any("登录" in str(row[2]) or "登录" in str(row[3]) for row in logins)

    # sections the pipeline does not collect are explicit placeholders
    assert workbook["暗网泄露"]["A2"].value == "未核查（无泄露情报库渠道）"
    assert workbook["APP资产"]["A2"].value == "未核查"
    assert workbook["漏洞汇总"]["B2"].value.startswith("本工具不做漏洞结论")


def test_machine_report_stays_the_primary_artifact(tmp_path):
    db = _seed(tmp_path)
    report = build_report(db, "t", tmp_path / "reports", include_xlsx=True)
    assert report["export"]["xlsx_status"] == "exported"
    assert (tmp_path / "reports" / "result.json").exists()
    assert (tmp_path / "reports" / "result.xlsx").exists()
    assert (tmp_path / "reports" / "analyst_report.xlsx").exists()
    # machine workbook does not gain the analyst sheets
    machine = load_workbook(tmp_path / "reports" / "result.xlsx")
    assert "域名资产" not in machine.sheetnames
