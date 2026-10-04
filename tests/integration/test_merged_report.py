"""C1: `report --merge` builds one cross-task delivery workbook with the
four main sheets, a per-task summary plus a total row, and a deterministic
output directory keyed by the sorted task-id set."""
import json

from openpyxl import load_workbook

from surface_intel.cli import main
from surface_intel.reporting.merged import merged_code
from surface_intel.state.database import Database


def _seed(tmp_path):
    db = Database(tmp_path / "task.sqlite")
    db.initialize()
    with db.connect() as connection:
        for tid in ("t1", "t2"):
            connection.execute(
                "INSERT INTO task(task_id,status,schema_version,scope_version) VALUES(?,'completed','1.0','scope-v1')",
                (tid,))
            connection.execute(
                """INSERT INTO asset(asset_id,task_id,asset_type,normalized_value,scope_status)
                   VALUES(?,?,'domain','example.com','confirmed')""", (f"a-{tid}", tid))
            payload = {"target": "www.example.com", "record_type": "A",
                       "records": ["203.0.113.5"], "status": "observed",
                       "observed_at": "2026-10-04T10:00:00+00:00"}
            connection.execute(
                "INSERT INTO observation(observation_id,task_id,source,target,payload_json) VALUES(?,?,?,?,?)",
                (f"o-{tid}", tid, "dns_resolver", "www.example.com", json.dumps(payload)))
    return db


def test_merged_report_layout_and_contents(tmp_path):
    _seed(tmp_path)
    code = main(["report", "--merge", "t2", "t1", "--task-root", str(tmp_path)])
    assert code == 0
    expected_dir = tmp_path / "reports" / "merged" / merged_code(["t1", "t2"])
    assert (expected_dir / "merged_report.xlsx").exists()
    manifest = json.loads((expected_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["task_ids"] == ["t1", "t2"]

    workbook = load_workbook(expected_dir / "merged_report.xlsx")
    assert workbook.sheetnames == ["报告摘要", "域名资产", "IP资产", "站点", "服务"]
    summary = [row for row in workbook["报告摘要"].iter_rows(values_only=True)][1:]
    assert [list(row)[:1] for row in summary] == [["t1"], ["t2"], ["合计"]]
    total = summary[-1]
    assert total[3] == 2 and total[4] == 2  # asset/observation totals
    domains = [row for row in workbook["域名资产"].iter_rows(values_only=True)][1:]
    assert {row[0] for row in domains} == {"t1", "t2"}
    assert domains[0][1] == "example.com"
    sites = [row for row in workbook["站点"].iter_rows(values_only=True)][1:]
    assert all(row[0] in {"t1", "t2"} for row in sites)


def test_merged_report_is_deterministic_per_task_set(tmp_path):
    _seed(tmp_path)
    assert main(["report", "--merge", "t1", "t2", "--task-root", str(tmp_path)]) == 0
    first = json.loads((tmp_path / "reports" / "merged" / merged_code(["t1", "t2"]) / "manifest.json").read_text(encoding="utf-8"))
    assert main(["report", "--merge", "t2", "t1", "--task-root", str(tmp_path)]) == 0
    second = json.loads((tmp_path / "reports" / "merged" / merged_code(["t1", "t2"]) / "manifest.json").read_text(encoding="utf-8"))
    # same set -> same directory, overwritten in place; only the timestamp moves
    first.pop("generated_at")
    second.pop("generated_at")
    assert first == second


def test_merge_and_positional_task_id_are_exclusive(tmp_path, capsys):
    _seed(tmp_path)
    code = main(["report", "t1", "--merge", "t2", "--task-root", str(tmp_path)])
    assert code == 2
    assert json.loads(capsys.readouterr().out)["error"] == "task_id_and_merge_exclusive"


def test_merge_unknown_task_refused(tmp_path, capsys):
    _seed(tmp_path)
    code = main(["report", "--merge", "t1", "ghost", "--task-root", str(tmp_path)])
    assert code == 2
    assert "task_not_found" in json.loads(capsys.readouterr().out)["error"]
