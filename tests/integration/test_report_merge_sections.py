"""R4.5: merged assets, relations and conflicts flow into the report."""
import json

from openpyxl import load_workbook

from surface_intel.merger import merge_task_observations
from surface_intel.policy.scope import Scope
from surface_intel.reporting.export import build_report
from surface_intel.state.database import Database


def _setup(tmp_path, scope):
    db = Database(tmp_path / "task.sqlite")
    db.initialize()
    with db.connect() as connection:
        connection.execute(
            "INSERT INTO task(task_id,status,schema_version,scope_version) VALUES('t','completed','1.0',?)",
            (scope.scope_version,),
        )
    return db


def _observation(db, task_id, source, payload, observed_at="2026-10-03T10:00:00+00:00"):
    payload = dict(payload)
    payload.setdefault("observed_at", "2026-10-03T10:00:00+00:00")
    with db.connect() as connection:
        connection.execute(
            """INSERT INTO observation(observation_id,task_id,source,target,payload_json)
               VALUES(?,?,?,?,?)""",
            (f"obs-{source}-{abs(hash(json.dumps(payload, sort_keys=True))) % 10**12}",
             task_id, source, str(payload.get("target", "")),
             json.dumps(payload, ensure_ascii=True, sort_keys=True)),
        )


def test_report_includes_merge_assets_relations_and_conflicts(tmp_path):
    scope = Scope(authorized_domains=("example.com", "*.example.com"),
                  authorized_ips=("203.0.113.5",),
                  confirmed=True, scope_version="scope-v1")
    db = _setup(tmp_path, scope)
    _observation(db, "t", "dns_resolver", {
        "target": "www.example.com", "record_type": "A", "records": ["203.0.113.5"],
        "status": "observed",
    })
    for org in ("Org A Ltd", "Org B Inc"):
        _observation(db, "t", "rdap_lookup", {
            "target": "203.0.113.5", "target_kind": "ip", "status": "observed",
            "record": {"organizations": [org]},
        }, observed_at=f"2026-10-03T1{('0' if org.startswith('Org A') else '1')}:00:00+00:00")
    merge_task_observations(db, "t", scope)

    output_dir = tmp_path / "report"
    report = build_report(db, "t", output_dir, include_xlsx=True)

    # JSON: dedicated merge sections + summary
    assert report["asset_relations"] and report["asset_conflicts"]
    summary = report["merge_summary"]
    assert summary["asset_count"] >= 3
    assert summary["confirmed_assets"] >= 2  # anchor domain + authorized IP
    assert summary["relation_count"] == len(report["asset_relations"])
    assert summary["unresolved_conflicts"] == 1
    assert "dns_resolver" in summary["sources"]

    # XLSX: dedicated sheets mirror the JSON sections
    workbook = load_workbook(output_dir / "result.xlsx")
    assert "架构关系" in workbook.sheetnames
    assert "归并冲突" in workbook.sheetnames
    relations_rows = list(workbook["架构关系"].iter_rows(values_only=True))
    assert relations_rows[0][0] == "relation_id"
    assert len(relations_rows) - 1 == len(report["asset_relations"])
    conflicts_rows = list(workbook["归并冲突"].iter_rows(values_only=True))
    assert conflicts_rows[0] == ("conflict_id", "asset_id", "attribute", "values_json", "status", "created_at")
    assert len(conflicts_rows) - 1 == summary["unresolved_conflicts"]
